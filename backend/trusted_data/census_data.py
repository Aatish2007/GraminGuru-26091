from openpyxl import load_workbook
from pathlib import Path


DATA_FOLDER = Path(__file__).parent

CENSUS_FILE_NAME = (
    "A-1_NO_OF_VILLAGES_TOWNS_HOUSEHOLDS_POPULATION_AND_AREA.xlsx"
)


# ============================================================
# LOCATION ALIASES
# ============================================================

LOCATION_ALIASES = {
    "kanpur": "Kanpur Nagar",
    "kanpur nagar": "Kanpur Nagar",

    "lucknow": "Lucknow",

    "varanasi": "Varanasi",
    "banaras": "Varanasi",

    "gorakhpur": "Gorakhpur",

    "agra": "Agra",

    "prayagraj": "Prayagraj",
    "allahabad": "Prayagraj",

    # Census 2011 uses Faizabad.
    # Modern user input may be Ayodhya.
    "ayodhya": "Faizabad",
    "faizabad": "Faizabad",

    "patna": "Patna",
    "ranchi": "Ranchi",
    "jaipur": "Jaipur",

    "mumbai": "Mumbai",
    "pune": "Pune",
    "nagpur": "Nagpur",

    "bhopal": "Bhopal",
    "indore": "Indore",

    "delhi": "Delhi"
}


# ============================================================
# NORMALIZE LOCATION
# ============================================================

def normalize_location(location):

    cleaned = " ".join(
        str(location).strip().split()
    )

    return LOCATION_ALIASES.get(
        cleaned.lower(),
        cleaned
    )


# ============================================================
# FIND CENSUS FILE
# ============================================================

def find_census_file():

    preferred_file = (
        DATA_FOLDER / CENSUS_FILE_NAME
    )

    if preferred_file.exists():
        return preferred_file

    excel_files = list(
        DATA_FOLDER.glob("*.xlsx")
    )

    if excel_files:
        return excel_files[0]

    return None


# ============================================================
# GET CENSUS DATA
# ============================================================

def get_census_data(location):

    requested_location = str(
        location or ""
    ).strip()

    normalized_location = normalize_location(
        requested_location
    )

    file_path = find_census_file()

    if not file_path:

        return {
            "success": False,
            "message": "Census Excel file not found."
        }

    workbook = None

    try:

        print(
            "\n[CENSUS] Requested:",
            requested_location
        )

        print(
            "[CENSUS] Normalized:",
            normalized_location
        )

        print(
            "[CENSUS] File:",
            file_path.name
        )

        workbook = load_workbook(
            file_path,
            read_only=True,
            data_only=True
        )

        # ----------------------------------------------------
        # Search every sheet
        # ----------------------------------------------------

        for sheet in workbook.worksheets:

            print(
                "[CENSUS] Checking sheet:",
                sheet.title
            )

            rows = sheet.iter_rows(
                values_only=True
            )

            # A-01 file has metadata/header rows.
            for _ in range(4):
                next(rows, None)

            # Header
            next(rows, None)

            for row in rows:

                if not row or len(row) < 15:
                    continue

                state_code = row[0]
                district_code = row[1]
                subdistrict_code = row[2]

                level = row[3]
                name = row[4]
                area_type = row[5]

                households = row[9]
                population = row[10]

                area = row[13]
                density = row[14]

                if not name:
                    continue

                name_clean = str(
                    name
                ).strip().lower()

                level_clean = str(
                    level or ""
                ).strip().lower()

                area_type_clean = str(
                    area_type or ""
                ).strip().lower()

                # ------------------------------------------------
                # IMPORTANT:
                # Only district + total rows.
                # ------------------------------------------------

                if name_clean != normalized_location.lower():
                    continue

                if level_clean != "district":
                    continue

                if area_type_clean != "total":
                    continue

                result = {

                    "success": True,

                    "location":
                        str(name).strip(),

                    "requested_location":
                        requested_location,

                    "normalized_location":
                        normalized_location,

                    "state_code":
                        state_code,

                    "district_code":
                        district_code,

                    "subdistrict_code":
                        subdistrict_code,

                    "households":
                        households,

                    "population":
                        population,

                    "area_km2":
                        area,

                    "population_density":
                        density,

                    "confidence":
                        "High",

                    "source":
                        "Census 2011 — Census of India",

                    "source_url":
                        "https://censusindia.gov.in/"
                }

                print(
                    "[CENSUS] FOUND:",
                    result
                )

                return result

        print(
            "[CENSUS] No district found:",
            normalized_location
        )

        return {

            "success": False,

            "message":
                f"No Census district found for '{requested_location}'.",

            "requested_location":
                requested_location,

            "normalized_location":
                normalized_location,

            "confidence":
                "Needs local verification",

            "source":
                "Census 2011 — Census of India"
        }

    except Exception as error:

        print(
            "[CENSUS] ERROR:",
            error
        )

        return {

            "success": False,

            "message":
                str(error),

            "requested_location":
                requested_location,

            "normalized_location":
                normalized_location
        }

    finally:

        if workbook is not None:

            try:
                workbook.close()
            except Exception:
                pass


# ============================================================
# TEST
# ============================================================

if __name__ == "__main__":

    print("\n--- KANPUR TEST ---")
    print(
        get_census_data("Kanpur")
    )

    print("\n--- AYODHYA TEST ---")
    print(
        get_census_data("Ayodhya")
    )